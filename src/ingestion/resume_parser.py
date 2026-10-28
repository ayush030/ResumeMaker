from pathlib import Path

import pymupdf
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_resume_text(file_path: Path) -> str:
    """Extract text from a PDF or DOCX resume."""

    file_path = Path(file_path)

    if not file_path.is_file():
        raise FileNotFoundError(
            f"Resume file not found: {file_path}"
        )

    if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported resume format. Use PDF or DOCX."
        )

    if file_path.suffix.lower() == ".pdf":
        text = _extract_pdf_text(file_path)
    else:
        text = _extract_docx_text(file_path)

    text = text.replace("\x00", "").strip()

    if len(text) < 50:
        raise ValueError(
            "Very little text could be extracted from the resume. "
            "The file may be scanned, image-based, or empty. "
            "OCR or manual review may be required."
        )

    return text


def _extract_pdf_text(file_path: Path) -> str:
    """Extract text from all PDF pages."""

    with pymupdf.open(file_path) as document:
        if document.is_encrypted:
            raise ValueError(
                "The PDF is encrypted. Provide an unlocked PDF."
            )

        if len(document) == 0:
            raise ValueError("The PDF contains no pages.")

        page_texts = [
            page.get_text(sort=True)
            for page in document
        ]

    return "\n\n".join(page_texts)


def _extract_docx_text(file_path: Path) -> str:
    """Extract paragraphs and table contents from DOCX."""

    document = Document(file_path)
    sections = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]
            if any(cells):
                sections.append(" | ".join(cells))

    return "\n".join(sections)
