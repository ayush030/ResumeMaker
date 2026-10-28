import tempfile
import unittest

from pathlib import Path
from docx import Document

from src.ingestion.resume_parser import extract_resume_text


class TestResumeParser(unittest.TestCase):

    def test_extract_docx_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "resume.docx"

            document = Document()
            document.add_paragraph(
                "Senior Backend Engineer with extensive experience."
            )
            document.add_paragraph(
                "Skills: Go, AWS, Kafka, Kubernetes."
            )
            document.save(path)

            text = extract_resume_text(path)

            self.assertIn("Senior Backend Engineer", text)
            self.assertIn("Go, AWS, Kafka", text)

    def test_missing_file(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "missing.pdf"

            with self.assertRaises(FileNotFoundError):
                extract_resume_text(path)

    def test_unsupported_extension(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "resume.txt"
            path.write_text("Some resume text")

            with self.assertRaises(ValueError):
                extract_resume_text(path)


if __name__ == "__main__":
    unittest.main()
